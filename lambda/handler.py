import base64
import html
import io
import json
import logging
import os
import re
from datetime import datetime, timezone
from email import policy as email_policy
from email.parser import BytesParser
from email.utils import parseaddr

import anthropic
import boto3
import openpyxl
import pypdf
from docx import Document

logger = logging.getLogger()
logger.setLevel(logging.INFO)

BUCKET_NAME = os.environ["BUCKET_NAME"]
NOTIFY_EMAIL = os.environ["NOTIFY_EMAIL"]
SENDER_EMAIL = os.environ["SENDER_EMAIL"]
SSM_API_KEY_PATH = os.environ["SSM_API_KEY_PATH"]
CLAUDE_MODEL = os.environ["CLAUDE_MODEL"]
# Only mail from these addresses is processed (comma-separated, case-insensitive).
ALLOWED_SENDERS = {
    a.strip().lower() for a in os.environ["ALLOWED_SENDERS"].split(",") if a.strip()
}

SYSTEM_PROMPT = """You are a document analysis assistant. The user will provide the body and attachments
from an email. Extract and summarize the key information from all documents provided.

The recipient usually forwards emails to you, so the outer sender is the recipient themselves.
Ignore the forwarding: never mention who forwarded the email. Describe the original message,
using the original sender from the forwarded headers in the body (e.g. "From: ...").

For each document, identify:
- Document type and purpose
- Key parties involved (names, companies, roles)
- Important dates and deadlines
- Financial figures or amounts
- Action items or required responses
- Any urgent or time-sensitive information

Return your response as a JSON object with these top-level keys:
- "summary": a 2-3 sentence plain-English summary of the email and all attachments. Begin it with
  "Email from <original sender's name or organization>", e.g. "Email from Lincoln Elementary about..."
- "documents": an array of objects, one per attachment, each with "filename", "type", and "key_points"
- "action_items": an array of objects for anything requiring the recipient's attention, each with
  "action" (a string, without the due date) and "due_date" (the deadline as "YYYY-MM-DD", or null if
  the item has no date). Resolve relative dates like "next Friday" against the date the original
  email was sent.
- "urgent": boolean, true if anything requires a response within 48 hours"""


# ── Attachment extraction helpers (no AWS calls) ───────────────────────────


def extract_pdf(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    return "\n".join(para.text for para in doc.paragraphs)


def extract_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join("" if v is None else str(v) for v in row)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def extract_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def image_to_content_block(data: bytes, content_type: str) -> dict:
    media_type = content_type.split(";")[0].strip()
    return {
        "type": "image",
        "source": {
            "type": "base64",
            "media_type": media_type,
            "data": base64.standard_b64encode(data).decode("utf-8"),
        },
    }


def extract_attachments(msg) -> list:
    """Return a list of Claude content blocks from all email attachments."""
    blocks = []
    for part in msg.walk():
        content_disposition = part.get_content_disposition()
        if content_disposition not in ("attachment", "inline"):
            continue

        filename = part.get_filename() or "unnamed"
        content_type = part.get_content_type()
        data = part.get_payload(decode=True)
        if not data:
            continue

        try:
            if content_type == "application/pdf":
                text = extract_pdf(data)
                blocks.append({"type": "text", "text": f"[Attachment: {filename}]\n{text}"})

            elif content_type in (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "application/msword",
            ):
                text = extract_docx(data)
                blocks.append({"type": "text", "text": f"[Attachment: {filename}]\n{text}"})

            elif content_type in (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel",
            ):
                text = extract_xlsx(data)
                blocks.append({"type": "text", "text": f"[Attachment: {filename}]\n{text}"})

            elif content_type in ("text/plain", "text/csv"):
                text = extract_text(data)
                blocks.append({"type": "text", "text": f"[Attachment: {filename}]\n{text}"})

            elif content_type in ("image/png", "image/jpeg"):
                blocks.append(image_to_content_block(data, content_type))

            else:
                logger.warning("Skipping unsupported attachment type: %s (%s)", filename, content_type)

        except Exception:
            logger.warning("Failed to extract attachment %s", filename, exc_info=True)

    return blocks


# ── AWS helpers ────────────────────────────────────────────────────────────


def get_ssm_parameter(name: str, ssm_client) -> str:
    response = ssm_client.get_parameter(Name=name, WithDecryption=True)
    return response["Parameter"]["Value"]


def save_to_s3(result: dict, message_id: str, s3_client) -> str:
    date_prefix = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    safe_id = message_id.strip("<>").replace("/", "_")
    key = f"parsed-output/{date_prefix}/{safe_id}.json"
    s3_client.put_object(
        Bucket=BUCKET_NAME,
        Key=key,
        Body=json.dumps(result, indent=2),
        ContentType="application/json",
    )
    return key


def send_email(subject: str, body_html: str, ses_client) -> None:
    ses_client.send_email(
        Source=SENDER_EMAIL,
        Destination={"ToAddresses": [NOTIFY_EMAIL]},
        Message={
            "Subject": {"Data": subject},
            "Body": {"Html": {"Data": body_html}},
        },
    )


def format_action_items(action_items: list) -> list:
    """Render action items as strings, dated ones first as "By Friday October 17: ..."."""
    dated, undated = [], []
    for item in action_items:
        if not isinstance(item, dict):
            undated.append(str(item))
            continue
        action = str(item.get("action", "")).strip()
        try:
            due = datetime.strptime(str(item.get("due_date")), "%Y-%m-%d").date()
        except ValueError:
            undated.append(action)
            continue
        # Weekday computed here, not by Claude, so it always matches the date.
        dated.append((due, f"By {due:%A %B} {due.day}: {action}"))
    return [text for _, text in sorted(dated, key=lambda d: d[0])] + undated


def build_email_body(result: dict, s3_key: str, rules_note: str = "") -> str:
    # Everything from Claude is derived from the forwarded email, which may be written
    # by anyone — escape it so it can't inject HTML (e.g. phishing links) into the summary.
    summary = html.escape(str(result.get("summary", "No summary available.")))
    action_items = format_action_items(result.get("action_items", []))
    action_html = (
        "<ul>" + "".join(f"<li>{html.escape(item)}</li>" for item in action_items) + "</ul>"
        if action_items
        else "<p>None</p>"
    )
    urgent_banner = (
        '<p style="color:red;font-weight:bold;">⚠ Urgent: response required within 48 hours.</p>'
        if result.get("urgent")
        else ""
    )
    return f"""
    <html><body>
    {urgent_banner}
    <h2>Email Parser Summary</h2>
    <p>{summary}</p>
    <h3>Action Items</h3>
    {action_html}
    <p style="color:grey;font-size:small;">{html.escape(rules_note)}<br>
    Full analysis saved to S3: {html.escape(s3_key)}</p>
    </body></html>
    """


def send_failure_notice(subject, message_id, key: str, exc: Exception, ses_client) -> None:
    """Tell NOTIFY_EMAIL that an email couldn't be processed. Never raises."""
    error = f"{type(exc).__name__}: {exc}"[:1000]
    body_html = f"""
    <html><body>
    <h2 style="color:#b00020;">Email Parser couldn't process an email</h2>
    <p><b>Subject:</b> {html.escape(str(subject))}<br>
    <b>Message-ID:</b> {html.escape(str(message_id))}<br>
    <b>Stored as:</b> {html.escape(key)}</p>
    <p><b>Error:</b> {html.escape(error)}</p>
    <p style="color:grey;font-size:small;">Full details are in CloudWatch Logs
    (log group /aws/lambda/email-parser). The original email stays in S3 for 30 days,
    so it can be reprocessed once the problem is fixed.</p>
    </body></html>
    """
    try:
        send_email(f"[Email Parser] Failed: {str(subject)[:150]}", body_html, ses_client)
        logger.info("Failure notice sent to %s", NOTIFY_EMAIL)
    except Exception:
        logger.exception("Could not send failure notice for email %s", message_id)


# ── Sender allowlist ───────────────────────────────────────────────────────


def sender_is_allowed(msg) -> bool:
    """Process only mail from an allowlisted address that SES authenticated.

    Anyone can email the parser address, and each processed email costs a Claude
    call. The From header alone is trivially forged, so also require SES's own
    DMARC verdict for the sender's domain (SPF or DKIM passed and aligned with the
    From domain). SES prepends its Authentication-Results header, so only the
    first one is trusted; any further down could have been written by the sender.
    """
    _, address = parseaddr(str(msg.get("From", "")))
    address = address.lower()
    if address not in ALLOWED_SENDERS:
        logger.warning("Dropping email from non-allowlisted sender: %s", address or "(none)")
        return False

    if str(msg.get("X-SES-Virus-Verdict", "")).strip().upper() == "FAIL":
        logger.warning("Dropping email from %s: SES virus scan failed", address)
        return False

    results = msg.get_all("Authentication-Results") or []
    ses_result = str(results[0]) if results else ""
    domain = re.escape(address.rsplit("@", 1)[-1])
    if not (
        ses_result.strip().lower().startswith("amazonses.com")
        and re.search(rf"\bdmarc=pass\b[^;]*\bheader\.from={domain}\b", ses_result, re.IGNORECASE)
    ):
        logger.warning("Dropping email from %s: SES authentication did not pass", address)
        return False

    return True


# ── Child rules ────────────────────────────────────────────────────────────

CHILD_RULES_KEY = "config/child-rules.json"
EMAIL_ADDRESS_RE = re.compile(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+")
# The "From:" line of a forwarded message's header block (Outlook, Gmail, Apple Mail).
FORWARDED_FROM_RE = re.compile(r"^[\s>*]*From:\**\s*(.*)$", re.IGNORECASE | re.MULTILINE)


def get_body_text(msg) -> str:
    """The message's own body (not attached emails), as plain text."""
    part = msg.get_body(preferencelist=("plain", "html"))
    if part is None:
        return ""
    text = part.get_content()
    if part.get_content_type() == "text/html":
        text = re.sub(r"(?is)<(style|script)\b.*?</\1>", "", text)
        text = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h\d)>", "\n", text)
        text = html.unescape(re.sub(r"<[^>]+>", "", text))
    return text


def find_attached_email(msg):
    """The first email attached to msg (Outlook's "Forward as Attachment"), or None."""
    for part in msg.walk():
        if part is not msg and part.get_content_type() == "message/rfc822":
            return part.get_content()
    return None


def find_original_sender(msg, body_text: str, attached):
    """The original sender's address, lowercased.

    Returns None if the email was forwarded but the forwarded header shows only a name.
    An email that wasn't forwarded at all is its own original.
    """
    if attached is not None:
        _, address = parseaddr(str(attached.get("From", "")))
        if address:
            return address.lower()
    match = FORWARDED_FROM_RE.search(body_text)
    if match:
        # Only the first "From:" line: later ones belong to older messages in the thread.
        address = EMAIL_ADDRESS_RE.search(match.group(1))
        return address.group(0).lower() if address else None
    _, address = parseaddr(str(msg.get("From", "")))
    return address.lower() or None


def load_child_rules(s3_client):
    """Return (children, error). Never raises: without rules, emails get a normal summary."""
    try:
        body = s3_client.get_object(Bucket=BUCKET_NAME, Key=CHILD_RULES_KEY)["Body"].read()
        children = json.loads(body).get("children", [])
        for child in children:
            if not isinstance(child.get("name"), str) or not child["name"].strip():
                raise ValueError("every child needs a \"name\"")
            for field in ("domains", "senders", "activities", "excluded_activities"):
                if not isinstance(child.get(field, []), list):
                    raise ValueError(f"{child['name']}: \"{field}\" must be a list")
            year = child.get("graduation_year")
            if year is not None and (type(year) is not int or not 2000 <= year <= 2100):
                raise ValueError(f"{child['name']}: \"graduation_year\" must be a year, e.g. 2030")
        return children, None
    except Exception as exc:
        logger.warning("Could not load child rules from %s", CHILD_RULES_KEY, exc_info=True)
        return [], f"{type(exc).__name__}: {exc}"[:300]


def match_children(children: list, address: str) -> list:
    """Children whose rules cover this sender: an exact sender, or a domain or its subdomain."""
    domain = address.rsplit("@", 1)[-1]
    matched = []
    for child in children:
        senders = {s.strip().lower() for s in child.get("senders", [])}
        domains = {d.strip().lower().lstrip("@") for d in child.get("domains", [])}
        if address in senders or any(domain == d or domain.endswith("." + d) for d in domains):
            matched.append(child)
    return matched


HIGH_SCHOOL_NAMES = {9: "freshman", 10: "sophomore", 11: "junior", 12: "senior"}


def grade_from_graduation_year(graduation_year: int, today) -> str:
    """e.g. 2030 → "9th grade (freshman, Class of 2030)" during the 2026-27 school year."""
    # The school year rolls over in August, so it ends next calendar year from August on.
    school_year_end = today.year + 1 if today.month >= 8 else today.year
    grade = 12 - (graduation_year - school_year_end)
    if grade > 12:
        return f"graduated (Class of {graduation_year})"
    if grade < 0:
        return "pre-kindergarten"
    if grade == 0:
        return "kindergarten"
    suffix = {1: "st", 2: "nd", 3: "rd"}.get(grade, "th")
    label = f"{grade}{suffix} grade"
    if grade in HIGH_SCHOOL_NAMES:
        return f"{label} ({HIGH_SCHOOL_NAMES[grade]}, Class of {graduation_year})"
    return f"{label} (Class of {graduation_year})"


def describe_child(child: dict, with_senders: bool = False) -> str:
    parts = [f"- {child['name']}"]
    # A hand-set grade wins (held back, unusual numbering); otherwise it advances each August.
    grade = child.get("grade")
    if not grade and child.get("graduation_year"):
        grade = grade_from_graduation_year(child["graduation_year"], datetime.now(timezone.utc).date())
    if grade:
        parts.append(f"grade: {grade}")
    if child.get("activities"):
        parts.append(f"activities: {', '.join(child['activities'])}")
    if child.get("excluded_activities"):
        parts.append(f"NOT part of: {', '.join(child['excluded_activities'])}")
    if child.get("notes"):
        parts.append(f"notes: {child['notes']}")
    if with_senders:
        parts.append(f"sender domains: {', '.join(child.get('domains', [])) or 'none'}")
        parts.append(f"sender addresses: {', '.join(child.get('senders', [])) or 'none'}")
    return "; ".join(parts)


FILTER_INSTRUCTIONS = """Keep only information that applies to the child or children it is about: their grade,
their activities, and anything for the whole school or family (closures, schedule changes, school-wide
events). Leave out information that is only for other grades. Always leave out anything about activities
the child is NOT part of. For activities in neither list, include them only if the child could take part
(e.g. open to their grade), and keep them brief. If nothing applies, say so in the summary and return no
action items. If the email is about more than one child, say which child each point is for, and start
each action item's "action" with the child's name."""


def child_rules_prompt(children: list, matched: list, address) -> str:
    """The system prompt section that applies child rules, or "" when none apply."""
    if matched:
        profiles = "\n".join(describe_child(c) for c in matched)
        return f"\n\nThis email is about:\n{profiles}\n\n{FILTER_INSTRUCTIONS}"
    if children and address is None:
        # The forwarded header shows only a name, so let Claude match the sender.
        profiles = "\n".join(describe_child(c, with_senders=True) for c in children)
        return (
            "\n\nThe original sender's email address couldn't be read from the forwarded header. "
            "These rules apply to email from particular sender domains or addresses:\n"
            f"{profiles}\n\n"
            "Apply a child's rule only if the email clearly comes from one of that child's sender "
            "domains or addresses (e.g. from the sender's signature or links). If one applies, "
            f"follow these instructions for that child:\n{FILTER_INSTRUCTIONS}\n\n"
            'Also return a "children" key: an array of the names of the children whose rules you '
            "applied, or an empty array if none."
        )
    return ""


# ── Lambda entry point ─────────────────────────────────────────────────────


def process_email(msg, message_id, subject, s3_client, ses_client, ssm_client) -> None:
    """Analyze an allowlisted email with Claude, save the result, and email a summary."""
    sender = msg.get("from", "unknown")
    body_text = get_body_text(msg)
    attached = find_attached_email(msg)

    content_blocks = [
        {
            "type": "text",
            "text": (
                f"Forwarded by: {sender}\nDate: {msg.get('date', 'unknown')}\n"
                f"Subject: {subject}\n\nBody:\n{body_text}"
            ),
        }
    ]
    if attached is not None:
        content_blocks.append({
            "type": "text",
            "text": (
                f"[Attached email]\nFrom: {attached.get('from', 'unknown')}\n"
                f"Date: {attached.get('date', 'unknown')}\n"
                f"Subject: {attached.get('subject', '(no subject)')}\n\n"
                f"Body:\n{get_body_text(attached)}"
            ),
        })
    content_blocks.extend(extract_attachments(msg))

    # Child rules: match the original sender, then tell Claude whose details to keep.
    original_sender = find_original_sender(msg, body_text, attached)
    children, rules_error = load_child_rules(s3_client)
    matched = match_children(children, original_sender) if original_sender else []
    names = [c["name"] for c in matched]
    if rules_error:
        rules_note = f"Child rules couldn't be loaded ({rules_error}), so nothing was filtered."
    elif matched:
        rules_note = f"Filtered for {', '.join(names)} (original sender: {original_sender})."
    elif original_sender:
        rules_note = f"No child rule matches the original sender ({original_sender}), so nothing was filtered."
    elif children:
        rules_note = "The original sender's address wasn't in the forwarded header, so Claude matched the rules."
    else:
        rules_note = ""
    logger.info("Original sender %s; matched children: %s", original_sender, names or "none")

    # Call Claude
    api_key = get_ssm_parameter(SSM_API_KEY_PATH, ssm_client)
    client = anthropic.Anthropic(api_key=api_key)
    claude_response = client.beta.messages.create(
        model=CLAUDE_MODEL,
        # Room for thinking plus the reply.
        max_tokens=16000,
        system=SYSTEM_PROMPT + child_rules_prompt(children, matched, original_sender),
        messages=[{"role": "user", "content": content_blocks}],
        # Summarizing and extracting from documents doesn't need deep reasoning.
        output_config={"effort": "low"},
        # If a safety classifier declines, retry server-side on another model.
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
    )

    if claude_response.stop_reason == "refusal":
        # Don't raise: Lambda would retry, and the retry would be declined too.
        details = claude_response.stop_details
        category = details.category if details else None
        logger.warning("Claude declined to analyze email %s (category: %s)", message_id, category)
        result = {
            "summary": f"Claude declined to analyze this email (category: {category}).",
            "documents": [], "action_items": [], "urgent": False,
        }
    else:
        # The response can begin with thinking blocks; the answer is in the text blocks.
        raw_text = "".join(b.text for b in claude_response.content if b.type == "text").strip()
        # Tolerate the JSON being wrapped in a Markdown code fence.
        if raw_text.startswith("```"):
            raw_text = raw_text.split("\n", 1)[-1].rsplit("```", 1)[0].strip()
        try:
            result = json.loads(raw_text)
        except json.JSONDecodeError:
            result = {"summary": raw_text, "documents": [], "action_items": [], "urgent": False}

    if not matched and original_sender is None:
        # Claude matched the rules; keep only names that are really in them.
        claimed = result.get("children") if isinstance(result.get("children"), list) else []
        names = [c["name"] for c in children if c["name"] in claimed]
        if names:
            rules_note = f"Filtered for {', '.join(names)} (sender matched by Claude)."
        elif children:
            rules_note = "The original sender's address wasn't found and no child rule matched, so nothing was filtered."

    result["_meta"] = {
        "message_id": message_id,
        "subject": subject,
        "sender": sender,
        "original_sender": original_sender,
        "children": names,
        "rules_note": rules_note,
        "processed_at": datetime.now(timezone.utc).isoformat(),
    }

    # Save to S3
    s3_key = save_to_s3(result, message_id, s3_client)
    logger.info("Saved analysis to %s", s3_key)

    # Send email summary
    child_tag = f"[{', '.join(names)}] " if names else ""
    email_subject = f"[Email Parser] {child_tag}{subject}"
    send_email(email_subject, build_email_body(result, s3_key, rules_note), ses_client)
    logger.info("Summary email sent to %s", NOTIFY_EMAIL)


def handler(event, context):
    s3_client = boto3.client("s3")
    ses_client = boto3.client("ses")
    ssm_client = boto3.client("ssm")

    record = event["Records"][0]["s3"]
    bucket = record["bucket"]["name"]
    key = record["object"]["key"]
    logger.info("Processing s3://%s/%s", bucket, key)

    # Used in the failure notice if the email itself can't be read.
    message_id, subject = key.split("/")[-1], "(email could not be read)"
    try:
        response = s3_client.get_object(Bucket=bucket, Key=key)
        raw_email = response["Body"].read()
        msg = BytesParser(policy=email_policy.default).parsebytes(raw_email)

        # Dropped mail gets no notice: anyone can email the parser address.
        if not sender_is_allowed(msg):
            return

        message_id = msg.get("Message-ID", message_id)
        subject = msg.get("subject", "(no subject)")
        process_email(msg, message_id, subject, s3_client, ses_client, ssm_client)
    except Exception as exc:
        # Notify instead of re-raising: a Lambda retry would repeat the paid Claude call
        # and send another notice, and the SDK already retries transient API errors.
        logger.exception("Failed to process email %s", message_id)
        send_failure_notice(subject, message_id, key, exc, ses_client)
